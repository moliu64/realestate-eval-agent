from __future__ import annotations

import warnings
from pathlib import Path

from langchain_core.documents import Document

SUPPORTED_TEXT = {".md", ".txt", ".csv"}
SUPPORTED_BINARY = {".pdf", ".docx"}
SUPPORTED_IMAGES = {".png", ".jpg", ".jpeg"}


def load_docs(docs_dir: str) -> list[Document]:
    root = Path(docs_dir)
    if not root.exists():
        raise FileNotFoundError(f"Docs directory does not exist: {root}")

    docs: list[Document] = []
    for path in sorted(root.iterdir()):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_TEXT | SUPPORTED_BINARY | SUPPORTED_IMAGES:
            warnings.warn(f"Skipping unsupported document format: {path.name}")
            continue
        document = load_document(path)
        if not document.page_content.strip():
            warnings.warn(f"Skipping empty document: {path.name}")
            continue
        docs.append(document)
    return docs


def load_document(path: Path) -> Document:
    suffix = path.suffix.lower()
    if suffix in SUPPORTED_TEXT:
        text = path.read_text(encoding="utf-8-sig")
    elif suffix == ".pdf":
        text = _load_pdf(path)
    elif suffix == ".docx":
        text = _load_docx(path)
    elif suffix in SUPPORTED_IMAGES:
        text = _load_image(path)
    else:
        raise ValueError(f"Unsupported document format: {suffix}")
    category = path.stem.split("_", 1)[0] if "_" in path.stem else "未分类"
    return Document(
        page_content=text.strip(),
        metadata={
            "source": path.name,
            "path": str(path),
            "file_type": suffix.lstrip("."),
            "category": category,
        },
    )


def _load_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF 入库需要安装 pypdf：pip install -r requirements.txt") from exc
    pages = []
    for index, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[第 {index} 页]\n{text.strip()}")
    return "\n\n".join(pages)


def _load_docx(path: Path) -> str:
    try:
        from docx import Document as WordDocument
    except ImportError as exc:
        raise RuntimeError("Word 入库需要安装 python-docx：pip install -r requirements.txt") from exc
    document = WordDocument(str(path))
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _load_image(path: Path) -> str:
    from app.vision import describe_image_file
    return describe_image_file(path)
